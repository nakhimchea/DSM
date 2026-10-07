"""Spreadsheet -> text serialisation (versioned runtime contract, ``TABLE_FORMAT_VERSION``).

The SAME function must be used to build training data and, at inference time, to
turn a user's ``.xlsx`` / ``.csv`` into the prompt - otherwise the model sees a
format it was not trained on. Applications can vendor this file.

Grid format (default), which exposes Excel addresses so formulas can be written::

    [Sheet "Loans" | A1:D4 | 3 data rows]
    |   | A | B | C | D |
    |---|---|---|---|---|
    | 1 | Branch | Officer | Amount (USD) | Status |
    | 2 | Phnom Penh | Dara | 12000 | Active |
    | 3 | Siem Reap | Lina | 8000 | Closed |
    | 4 | Phnom Penh | Sok | 5500 | Active |

Formula cells are written as ``=SUM(C2:C4) → 25500`` (formula, then cached
value). Merged ranges are listed after the grid; dates are ISO 8601; numbers are
raw values (no thousands separators); long sheets are truncated with an
explicit note pointing to the ``table_query`` tool.
"""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path
from typing import Any

from dsm.tables.formula import index_to_col

TABLE_FORMAT_VERSION = "1"
MAX_CELL_CHARS = 120


def format_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if value != value:  # NaN
            return ""
        return str(int(value)) if value.is_integer() else f"{value:.10g}"
    if isinstance(value, (dt.datetime, dt.date)):
        if isinstance(value, dt.datetime) and (value.hour, value.minute, value.second) == (0, 0, 0):
            return value.date().isoformat()
        return value.isoformat()
    text = str(value).replace("\n", " ").replace("|", "\\|").strip()
    return text if len(text) <= MAX_CELL_CHARS else text[: MAX_CELL_CHARS - 1] + "…"


def grid_to_text(
    rows: list[list[Any]],
    sheet: str = "Sheet1",
    *,
    first_row: int = 1,
    first_col: int = 1,
    header_rows: int = 1,
    max_rows: int = 60,
    merged: list[str] | None = None,
    style: str = "grid",
) -> str:
    """Serialise a 2-D list of cell values (top-left cell at ``first_row``/``first_col``)."""
    n_rows = len(rows)
    n_cols = max((len(r) for r in rows), default=0)
    last_col = index_to_col(first_col + max(n_cols, 1) - 1)
    span = f"{index_to_col(first_col)}{first_row}:{last_col}{first_row + max(n_rows, 1) - 1}"
    data_rows = max(n_rows - header_rows, 0)
    lines = [f'[Sheet "{sheet}" | {span} | {data_rows} data rows]']
    shown = rows[: header_rows + max_rows]
    if style == "grid":
        letters = [index_to_col(first_col + i) for i in range(n_cols)]
        lines.append("|   | " + " | ".join(letters) + " |")
        lines.append("|---|" + "---|" * n_cols)
        for offset, row in enumerate(shown):
            cells = [format_cell(v) for v in row] + [""] * (n_cols - len(row))
            lines.append(f"| {first_row + offset} | " + " | ".join(cells) + " |")
    elif style == "markdown":
        header = [format_cell(v) for v in (shown[0] if shown else [])]
        lines.append("| " + " | ".join(header) + " |")
        lines.append("|" + "---|" * len(header))
        for row in shown[1:]:
            cells = [format_cell(v) for v in row] + [""] * (len(header) - len(row))
            lines.append("| " + " | ".join(cells) + " |")
    else:
        raise ValueError("style must be 'grid' or 'markdown'")
    if n_rows > len(shown):
        lines.append(
            f"… {n_rows - len(shown)} more rows not shown (use the table_query tool for the full sheet)"
        )
    if merged:
        lines.append("Merged cells: " + ", ".join(merged))
    return "\n".join(lines)


def dataframe_to_text(df, sheet: str = "Sheet1", **kwargs) -> str:
    """A pandas DataFrame as a sheet whose row 1 holds the column names."""
    rows = [list(df.columns)] + df.astype(object).where(df.notna(), None).values.tolist()
    return grid_to_text(rows, sheet, **kwargs)


def workbook_to_text(path: str | Path, max_rows: int = 60, style: str = "grid") -> str:
    """Serialise every sheet of an ``.xlsx`` (formulas with cached values) or a ``.csv`` file."""
    path = Path(path)
    if path.suffix.lower() in (".csv", ".tsv"):
        import csv

        delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
        with open(path, newline="", encoding="utf-8-sig") as handle:
            rows = [[_parse_number(c) for c in row] for row in csv.reader(handle, delimiter=delimiter)]
        return grid_to_text(rows, path.stem, max_rows=max_rows, style=style)

    import openpyxl

    formulas = openpyxl.load_workbook(path, data_only=False)
    values = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for ws in formulas.worksheets:
        cached = values[ws.title]
        min_row, min_col, max_row, max_col = ws.min_row, ws.min_column, ws.max_row, ws.max_column
        merged = [str(r) for r in ws.merged_cells.ranges]
        merged_fill = {}
        for rng in ws.merged_cells.ranges:
            top_left = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    merged_fill[(r, c)] = top_left
        rows = []
        for r in range(min_row, max_row + 1):
            row = []
            for c in range(min_col, max_col + 1):
                value = merged_fill.get((r, c), ws.cell(r, c).value)
                if isinstance(value, str) and value.startswith("="):
                    result = cached.cell(r, c).value
                    value = f"{value} → {format_cell(result)}" if result is not None else value
                row.append(value)
            rows.append(row)
        parts.append(
            grid_to_text(
                rows,
                ws.title,
                first_row=min_row,
                first_col=min_col,
                max_rows=max_rows,
                merged=merged,
                style=style,
            )
        )
    return "\n\n".join(parts)


def _parse_number(text: str) -> Any:
    stripped = text.strip()
    if not stripped:
        return None
    try:
        number = (
            float(stripped.replace(",", ""))
            if stripped.replace(",", "").replace(".", "", 1).lstrip("-").isdigit()
            else None
        )
    except ValueError:
        number = None
    if number is None:
        return stripped
    return int(number) if number.is_integer() and "." not in stripped else number


def cells_from_rows(rows: list[list[Any]], first_row: int = 1, first_col: int = 1) -> dict[str, Any]:
    """``{"A1": value}`` map for the formula evaluator."""
    cells = {}
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            cells[f"{index_to_col(first_col + c)}{first_row + r}"] = value
    return cells


_GRID_HEADER = re.compile(r'^\[Sheet "(?P<sheet>.*)" \| (?P<span>[A-Z]+\d+:[A-Z]+\d+) \|')


def _cell_value(text: str) -> Any:
    text = text.strip().replace("\\|", "|")
    if text == "":
        return None
    if text in ("TRUE", "FALSE"):
        return text == "TRUE"
    try:
        number = float(text)
    except ValueError:
        return text
    return int(number) if number.is_integer() and "." not in text and "e" not in text.lower() else number


def parse_grid(text: str) -> tuple[str, dict[str, Any]]:
    """Inverse of the ``grid`` style: ``(sheet name, {"A1": value})`` (formula cells keep their text).

    Used by tests and evaluation to re-check answers against exactly what the
    model was shown.
    """
    lines = text.strip().splitlines()
    match = _GRID_HEADER.match(lines[0]) if lines else None
    if not match:
        raise ValueError("not a grid-serialised sheet")
    letters = [c.strip() for c in re.split(r"(?<!\\)\|", lines[1].strip().strip("|"))][1:]
    cells: dict[str, Any] = {}
    for line in lines[3:]:
        if not line.startswith("| "):
            break
        parts = re.split(r"(?<!\\)\|", line.strip()[1:-1])
        row = parts[0].strip()
        for letter, raw in zip(letters, parts[1:], strict=False):
            if (value := _cell_value(raw)) is not None:
                cells[f"{letter}{row}"] = value
    return match.group("sheet"), cells
